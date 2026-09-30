from docx import Document
from docx.shared import Pt, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

doc = Document()
section = doc.sections[0]
section.page_width, section.page_height = section.page_height, section.page_width
section.left_margin = section.right_margin = Cm(1.5)

C_SALMON = 'FCB6A5'  # \cellcolor{color_blue} RGB(252,182,165)
C_LBLUE  = 'EBEBFF'  # \cellcolor{blue!8}

def set_cell_bg(cell, hex_color):
    tc = cell._tc
    tcPr = tc.find(qn('w:tcPr'))
    if tcPr is None:
        tcPr = OxmlElement('w:tcPr')
        tc.insert(0, tcPr)
    for old in tcPr.findall(qn('w:shd')):
        tcPr.remove(old)
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'), hex_color)
    tcPr.append(shd)

def cell_text(cell, text, bold=False):
    para = cell.paragraphs[0]
    para.clear()
    para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = para.add_run(text)
    run.font.name = 'Times New Roman'
    run.font.size = Pt(10)
    run.font.bold = bold

def _get_tcBorders(tc_el):
    tcPr = tc_el.find(qn('w:tcPr'))
    if tcPr is None:
        tcPr = OxmlElement('w:tcPr')
        tc_el.insert(0, tcPr)
    tcB = tcPr.find(qn('w:tcBorders'))
    if tcB is None:
        tcB = OxmlElement('w:tcBorders')
        tcPr.append(tcB)
    return tcB

def border_on(tc_el, side, sz, color='000000'):
    tcB = _get_tcBorders(tc_el)
    for old in tcB.findall(qn(f'w:{side}')): tcB.remove(old)
    el = OxmlElement(f'w:{side}')
    el.set(qn('w:val'), 'single')
    el.set(qn('w:sz'), str(sz))
    el.set(qn('w:color'), color)
    el.set(qn('w:space'), '0')
    tcB.append(el)

def border_off(tc_el, side):
    tcB = _get_tcBorders(tc_el)
    for old in tcB.findall(qn(f'w:{side}')): tcB.remove(old)
    el = OxmlElement(f'w:{side}')
    el.set(qn('w:val'), 'none')
    el.set(qn('w:sz'), '0')
    el.set(qn('w:color'), 'auto')
    el.set(qn('w:space'), '0')
    tcB.append(el)

# 列: col0=Methods, col1-3=4-view, col4-6=6-view, col7-9=9-view
NCOLS, NROWS = 10, 9
table = doc.add_table(rows=NROWS, cols=NCOLS)
table.style = 'Table Grid'

# 清除表级边框
tbl = table._tbl
tblPr = tbl.find(qn('w:tblPr'))
if tblPr is None:
    tblPr = OxmlElement('w:tblPr'); tbl.insert(0, tblPr)
tblBdr = tblPr.find(qn('w:tblBorders'))
if tblBdr is None:
    tblBdr = OxmlElement('w:tblBorders'); tblPr.append(tblBdr)
for side in ['top','left','bottom','right','insideH','insideV']:
    for old in tblBdr.findall(qn(f'w:{side}')): tblBdr.remove(old)
    el = OxmlElement(f'w:{side}')
    el.set(qn('w:val'),'none'); el.set(qn('w:sz'),'0'); el.set(qn('w:color'),'auto')
    tblBdr.append(el)

# ── 表头 ──
table.cell(0, 0).merge(table.cell(1, 0))
cell_text(table.cell(0, 0), 'Methods', bold=True)

for start, label in [(1,'MipNeRF360 (4-view)'), (4,'MipNeRF360 (6-view)'), (7,'MipNeRF360 (9-view)')]:
    c = table.cell(0, start); c.merge(table.cell(0, start+2))
    cell_text(c, label, bold=True)

metrics = ['LPIPS*↓', 'PSNR↑', 'SSIM↑']
for i, m in enumerate(metrics):
    cell_text(table.cell(1, 1+i), m, bold=True)
    cell_text(table.cell(1, 4+i), m, bold=True)
    cell_text(table.cell(1, 7+i), m, bold=True)

# ── 数据 ──
rows_data = [
    ('RegNeRF',      '19.88','12.59','0.841', '20.72','13.41','0.847', '19.70','13.68','0.852'),
    ('SparseNeRF',   '17.76','12.83','0.845', '19.74','13.42','0.901', '21.56','14.36','0.823'),
    ('3DGS',         '10.80','20.31','0.899',  '8.38','22.12','0.913',  '6.42','24.29','0.930'),
    ('CoR-GS',       '11.76','20.45','0.712',  '9.28','22.57','0.886',  '7.04','24.73','0.916'),
    ('DropGaussian', '11.57','21.76','0.837',  '8.61','22.98','0.895',  '8.02','24.10','0.915'),
    ('D²GS',          '2.58','22.35','0.923',  '2.34','24.74','0.937',  '2.11','27.13','0.941'),
    ('Ours',          '3.89','25.91','0.948',  '3.43','27.93','0.953',  '3.02','29.21','0.961'),
]

# 精确对照 LaTeX 逐格高亮
# r_idx: 0=RegNeRF .. 6=Ours; col: 1-9
highlights = {
    # D²GS (r=5): LPIPS全三组→color_blue; PSNR/SSIM全三组→blue!8
    (5,1):C_SALMON, (5,2):C_LBLUE,  (5,3):C_LBLUE,
    (5,4):C_SALMON, (5,5):C_LBLUE,  (5,6):C_LBLUE,
    (5,7):C_SALMON, (5,8):C_LBLUE,  (5,9):C_LBLUE,
    # Ours (r=6): LPIPS全三组→blue!8; PSNR/SSIM全三组→color_blue
    (6,1):C_LBLUE,  (6,2):C_SALMON, (6,3):C_SALMON,
    (6,4):C_LBLUE,  (6,5):C_SALMON, (6,6):C_SALMON,
    (6,7):C_LBLUE,  (6,8):C_SALMON, (6,9):C_SALMON,
}

for r_idx, row_data in enumerate(rows_data):
    tr = r_idx + 2
    is_ours = (row_data[0] == 'Ours')
    cell_text(table.cell(tr, 0), row_data[0], bold=is_ours)
    for ci, val in enumerate(row_data[1:]):
        col = ci + 1
        cell = table.cell(tr, col)
        cell_text(cell, val, bold=is_ours)
        if (r_idx, col) in highlights:
            set_cell_bg(cell, highlights[(r_idx, col)])

# ── 列宽 ──
col_widths = [Cm(2.8)] + [Cm(1.9)] * 9
for row in table.rows:
    for i, cell in enumerate(row.cells):
        if i < len(col_widths):
            cell.width = col_widths[i]

# ── 边框（同FFLL：竖线在col0右/col3右/col6右；横线仅顶/表头后/底）──
THICK = 18
VTHIN = 6
RIGHT_AFTER = {0, 3, 6}

for row_idx, tr_el in enumerate(table._tbl.findall(qn('w:tr'))):
    col_idx = 0
    for tc_el in tr_el.findall(qn('w:tc')):
        tcPr_el = tc_el.find(qn('w:tcPr'))
        grid_span = 1
        if tcPr_el is not None:
            gs = tcPr_el.find(qn('w:gridSpan'))
            if gs is not None:
                grid_span = int(gs.get(qn('w:val'), '1'))
        last_col = col_idx + grid_span - 1

        if row_idx == 0:
            border_on(tc_el, 'top', THICK)
        else:
            border_off(tc_el, 'top')

        if row_idx == 1 or row_idx == NROWS - 1:
            border_on(tc_el, 'bottom', THICK)
        else:
            border_off(tc_el, 'bottom')

        border_off(tc_el, 'left')

        if last_col in RIGHT_AFTER:
            border_on(tc_el, 'right', VTHIN)
        else:
            border_off(tc_el, 'right')

        col_idx += grid_span

output = '/Users/zhangao/Desktop/中国传媒大学/论文/（Best）journal of electronic imaging/latex/MipNeRF360数据集/table.docx'
doc.save(output)
print(f'成功保存: {output}')
