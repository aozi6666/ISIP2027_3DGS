from docx import Document
from docx.shared import Pt, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

doc = Document()
section = doc.sections[0]
section.left_margin = section.right_margin = Cm(2.0)

def cell_text(cell, text, bold=False):
    para = cell.paragraphs[0]
    para.clear()
    para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = para.add_run(text)
    run.font.name = 'Times New Roman'
    run.font.size = Pt(10)
    run.font.bold = bold

def _get_tcBorders(tc_el):
    tcPr = tc_el.find(qn('w:tcPr'))
    if tcPr is None:
        tcPr = OxmlElement('w:tcPr')
        tc_el.insert(0, tcPr)
    tcB = tcPr.find(qn('w:tcBorders'))
    if tcB is None:
        tcB = OxmlElement('w:tcBorders')
        tcPr.append(tcB)
    return tcB

def border_on(tc_el, side, sz, color='000000'):
    tcB = _get_tcBorders(tc_el)
    for old in tcB.findall(qn(f'w:{side}')): tcB.remove(old)
    el = OxmlElement(f'w:{side}')
    el.set(qn('w:val'), 'single')
    el.set(qn('w:sz'), str(sz))
    el.set(qn('w:color'), color)
    el.set(qn('w:space'), '0')
    tcB.append(el)

def border_off(tc_el, side):
    tcB = _get_tcBorders(tc_el)
    for old in tcB.findall(qn(f'w:{side}')): tcB.remove(old)
    el = OxmlElement(f'w:{side}')
    el.set(qn('w:val'), 'none')
    el.set(qn('w:sz'), '0')
    el.set(qn('w:color'), 'auto')
    el.set(qn('w:space'), '0')
    tcB.append(el)

# 6列，4行（1行表头 + 3行数据）
NCOLS, NROWS = 6, 4
table = doc.add_table(rows=NROWS, cols=NCOLS)
table.style = 'Table Grid'

# 清除表级边框
tbl = table._tbl
tblPr = tbl.find(qn('w:tblPr'))
if tblPr is None:
    tblPr = OxmlElement('w:tblPr'); tbl.insert(0, tblPr)
tblBdr = tblPr.find(qn('w:tblBorders'))
if tblBdr is None:
    tblBdr = OxmlElement('w:tblBorders'); tblPr.append(tblBdr)
for side in ['top','left','bottom','right','insideH','insideV']:
    for old in tblBdr.findall(qn(f'w:{side}')): tblBdr.remove(old)
    el = OxmlElement(f'w:{side}')
    el.set(qn('w:val'),'none'); el.set(qn('w:sz'),'0'); el.set(qn('w:color'),'auto')
    tblBdr.append(el)

# ── 表头 ──
headers = ['Visual Hull Init', 'Depth Densification', 'Visibility Balancing', 'SSIM↑', 'PSNR↑', 'LPIPS↓']
for i, h in enumerate(headers):
    cell_text(table.cell(0, i), h, bold=True)

# ── 数据 ──
# √ 用 ✓ 表示，空格表示无
rows_data = [
    ('✓', '',  '',  '0.865', '17.51', '1.128'),
    ('✓', '✓', '',  '0.902', '24.96', '0.879'),
    ('✓', '✓', '✓', '0.943', '26.18', '0.609'),
]

for r_idx, row_data in enumerate(rows_data):
    tr = r_idx + 1
    is_last = (r_idx == 2)
    for ci, val in enumerate(row_data):
        cell_text(table.cell(tr, ci), val, bold=is_last)

# ── 列宽 ──
col_widths = [Cm(3.2), Cm(3.8), Cm(3.8), Cm(2.0), Cm(2.0), Cm(2.0)]
for row in table.rows:
    for i, cell in enumerate(row.cells):
        cell.width = col_widths[i]

# ── 边框 ──
# \toprule 表顶（粗）
# \midrule 表头后（细）
# \bottomrule 表底（粗）
# 无竖线，无行间线
THICK = 18  # \toprule / \bottomrule
THIN  = 6   # \midrule

for row_idx, tr_el in enumerate(table._tbl.findall(qn('w:tr'))):
    for tc_el in tr_el.findall(qn('w:tc')):
        # 上边框
        if row_idx == 0:
            border_on(tc_el, 'top', THICK)
        else:
            border_off(tc_el, 'top')

        # 下边框
        if row_idx == 0:
            border_on(tc_el, 'bottom', THIN)   # \midrule
        elif row_idx == NROWS - 1:
            border_on(tc_el, 'bottom', THICK)  # \bottomrule
        else:
            border_off(tc_el, 'bottom')

        border_off(tc_el, 'left')
        border_off(tc_el, 'right')

output = '/Users/zhangao/Desktop/中国传媒大学/论文/（Best）journal of electronic imaging/latex/消融实验/table.docx'
doc.save(output)
print(f'成功保存: {output}')
